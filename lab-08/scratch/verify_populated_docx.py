import docx

doc = docx.Document('Lab 08 Worksheet .docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    first_cell = table.rows[0].cells[0].text.strip().replace('\n', ' ')[:70]
    blips = table._element.xpath('.//*[local-name()="blip"]')
    print(f"Table {t_idx:2d} ({rows}x{cols}): text='{first_cell}' | images={len(blips)}")

# Print paragraphs
print("\n--- Paragraph summary ---")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print(f"P{i:2d}: {t[:60]}")
